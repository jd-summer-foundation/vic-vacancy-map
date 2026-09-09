"""
Convert the master "Disability Provider Vacancies.xlsx" workbook's "Vacancies"
sheet into a flat all_rows.json (list of {column: value} dicts) -- the input
build_data.py expects.

Usage:
    python3 xlsx_to_json.py "Disability Provider Vacancies.xlsx"

Writes ./all_rows.json (gitignored -- regenerate it locally, don't commit it;
it's a full export of every provider's listing text, not just the VIC subset
this map needs).
"""
import json
import sys
from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else "Disability Provider Vacancies.xlsx"

wb = load_workbook(path, read_only=True, data_only=True)
ws = wb["Vacancies"]

rows_iter = ws.iter_rows(values_only=True)
headers = next(rows_iter)

rows = []
for r in rows_iter:
    row = {h: v for h, v in zip(headers, r) if h is not None}
    # openpyxl gives None for empty cells; normalise to '' to match the
    # scrape's own output (build_data.py checks truthiness either way, but
    # this keeps a diff against the scrape's own all_rows.json clean)
    row = {k: ("" if v is None else v) for k, v in row.items()}
    rows.append(row)

with open("all_rows.json", "w") as f:
    json.dump(rows, f, ensure_ascii=False)

print(f"wrote all_rows.json: {len(rows)} rows, {len(headers)} columns")
